"""골드 시트 읽기와 '골드 정답 반영'용 문장 → 정답 맵.

사람이 정답을 확정한 문장(골드)이 새 실행에 다시 나오면 모델 판정 대신 골드 정답을 그대로 쓴다.
비교는 공백을 정리한 문장 전체가 같을 때만. 여러 골드에 같은 문장이 다른 정답으로 있으면 가장 최근에 고친 파일이 이긴다."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .models import GRADES

GOLD_HEADERS = ["번호", "파일", "페이지", "제목", "줄", "문장", "모델 라벨", "후처리 라벨", "판정", "확신", "점수", "가중치", "정답", "메모"]
LABELS = list(GRADES)
GOLD_SUFFIX = "_gold.xlsx"


def norm_label(label) -> str:
    s = str(label or "").strip().upper()
    return "X" if s in ("X?", "X") else s


def norm_text(text) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def load_gold(path: Path) -> list[dict]:
    """'라벨링' 시트에서 정답이 A~D/X 인 행만. 열은 이름으로 찾으므로 옛 골드도 읽힌다."""
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        ws = wb["라벨링"]
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(rows_iter)]
        idx = {h: i for i, h in enumerate(header)}
        rows = []
        for r in rows_iter:
            if r is None or idx.get("문장") is None or r[idx["문장"]] is None:
                continue
            gold = norm_label(r[idx["정답"]]) if "정답" in idx else ""
            if gold not in LABELS:
                continue
            rows.append({h: r[idx[h]] for h in GOLD_HEADERS if h in idx} | {"정답": gold})
        return rows
    finally:
        wb.close()


def find_gold_files(paths: list[str | Path], base_dir: Path | None = None) -> list[Path]:
    """파일이면 그대로, 폴더면 그 안의 *_gold.xlsx (하위 폴더 포함). 없는 경로는 조용히 건너뛴다."""
    out: list[Path] = []
    for p in paths:
        p = Path(p).expanduser()
        if not p.is_absolute() and base_dir is not None:
            p = base_dir / p
        if p.is_dir():
            out.extend(sorted(f for f in p.rglob(f"*{GOLD_SUFFIX}") if not f.name.startswith("~$")))
        elif p.is_file():
            out.append(p)
    seen: set[Path] = set()
    uniq = []
    for f in out:
        if f.resolve() not in seen:
            seen.add(f.resolve())
            uniq.append(f)
    return uniq


@dataclass
class GoldMap:
    labels: dict[str, str] = field(default_factory=dict)   # 정리한 문장 → 정답
    files: list[Path] = field(default_factory=list)
    rows: int = 0
    conflicts: int = 0                                      # 같은 문장, 다른 정답 (최근 파일이 이김)

    @classmethod
    def load(cls, paths: list[str | Path], base_dir: Path | None = None) -> "GoldMap":
        gm = cls()
        files = find_gold_files(paths, base_dir)
        for f in sorted(files, key=lambda p: p.stat().st_mtime):      # 오래된 것부터 → 최근 파일이 덮어씀
            try:
                rows = load_gold(f)
            except Exception:
                continue
            gm.files.append(f)
            for r in rows:
                key = norm_text(r.get("문장"))
                if not key:
                    continue
                gm.rows += 1
                old = gm.labels.get(key)
                if old is not None and old != r["정답"]:
                    gm.conflicts += 1
                gm.labels[key] = r["정답"]
        return gm

    def get(self, text: str) -> str | None:
        return self.labels.get(norm_text(text))

    def __len__(self) -> int:
        return len(self.labels)

    def __bool__(self) -> bool:
        return bool(self.labels)

    def describe(self) -> str:
        if not self.files:
            return "없음"
        return f"{len(self.files)}개 파일, {len(self.labels)}문장" + (f", 충돌 {self.conflicts}" if self.conflicts else "")

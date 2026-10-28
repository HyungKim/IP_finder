"""PPTX 리더 (python-pptx). 제목 플레이스홀더, 본문 문단(계층 유지), 표, 그룹 도형 안의 글, SmartArt, 차트 제목을 읽는다.
발표자 노트는 읽지 않는다. 숨긴 슬라이드도 읽는다(보고 자료에서는 뺀 쪽에 기술 내용이 남아 있는 일이 많다).

- 그룹 도형은 그룹 변환을 풀어 슬라이드 절대 좌표로 바꾼 뒤, 모든 도형을 위→아래, 같은 행(세로로 절반 이상 겹침)에서는
  왼쪽→오른쪽 순서로 읽는다. XML 순서(만든 순서)는 사람이 읽는 순서와 다를 수 있다.
- 제목 플레이스홀더가 없거나 비어 있으면 슬라이드 위쪽(높이 35% 안)에서 글씨가 가장 큰 텍스트 상자를 제목으로 본다.
- 표는 행마다 '열 제목: 셀 | …' 한 줄(source="table"), SmartArt 와 차트 제목은 source="diagram" 줄로 만든다.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.group import GroupShape

from ..models import Document, Line, Page
from ._text import clean, table_lines

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
OLE_URI = "http://schemas.openxmlformats.org/presentationml/2006/ole"

_GRAPHIC_TYPES = {
    MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE,
    MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.LINKED_PICTURE, MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.DIAGRAM,
}
TITLE_MIN_RATIO = 0.8      # 제목 후보는 슬라이드 본문 글씨의 이 배수 이상 (작은 머리글 상자 제외)
TITLE_ZONE = 0.35      # 슬라이드 높이의 이 비율 안에 있어야 제목 후보
TITLE_MAX_CHARS = 80   # 이보다 길면 제목이 아니라 본문

Transform = Callable[[float, float, float, float], tuple[float, float, float, float]]


def _identity(x: float, y: float, w: float, h: float) -> tuple[float, float, float, float]:
    return x, y, w, h


def group_transform(group_element: Any, parent: Transform) -> Transform:
    """그룹 도형의 자식 좌표를 슬라이드 좌표로 바꾸는 변환 (회전·뒤집기는 반영하지 않음)."""
    xfrm = None
    group_props = group_element.find(P + "grpSpPr")
    if group_props is not None:
        xfrm = group_props.find(A + "xfrm")
    if xfrm is None:
        return parent
    off, ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
    ch_off, ch_ext = xfrm.find(A + "chOff"), xfrm.find(A + "chExt")
    if off is None or ext is None or ch_off is None or ch_ext is None:
        return parent
    ox, oy = float(off.get("x", 0)), float(off.get("y", 0))
    cx, cy = float(ext.get("cx", 0)), float(ext.get("cy", 0))
    cox, coy = float(ch_off.get("x", 0)), float(ch_off.get("y", 0))
    ccx, ccy = float(ch_ext.get("cx", 0)), float(ch_ext.get("cy", 0))
    sx = cx / ccx if ccx else 1.0
    sy = cy / ccy if ccy else 1.0

    def transform(x: float, y: float, w: float, h: float) -> tuple[float, float, float, float]:
        return parent(ox + (x - cox) * sx, oy + (y - coy) * sy, w * sx, h * sy)

    return transform


@dataclass
class ShapeInfo:
    shape: Any
    left: float
    top: float
    width: float
    height: float
    xml_order: int
    in_group: bool = False
    reading_order: int = 0


def collect_shapes(shapes: Any, transform: Transform = _identity, in_group: bool = False,
                   out: list[ShapeInfo] | None = None) -> list[ShapeInfo]:
    """그룹을 풀어 모든 말단 도형을 슬라이드 절대 좌표와 함께 모은다."""
    out = [] if out is None else out
    for shape in shapes:
        try:
            box = transform(float(shape.left or 0), float(shape.top or 0), float(shape.width or 0), float(shape.height or 0))
        except Exception:
            box = (0.0, 0.0, 0.0, 0.0)
        if isinstance(shape, GroupShape):
            collect_shapes(shape.shapes, group_transform(shape._element, transform), True, out)
        else:
            out.append(ShapeInfo(shape, *box, xml_order=len(out), in_group=in_group))
    return out


def assign_reading_order(infos: list[ShapeInfo]) -> list[ShapeInfo]:
    """위→아래, 같은 행에서는 왼쪽→오른쪽. 세로로 절반 이상 겹치면 같은 행으로 본다."""
    rows: list[list[ShapeInfo]] = []
    for info in sorted(infos, key=lambda i: (i.top, i.left)):
        placed = False
        for row in rows:
            anchor = row[0]
            overlap = min(anchor.top + anchor.height, info.top + info.height) - max(anchor.top, info.top)
            base = min(anchor.height, info.height)
            if base > 0 and overlap / base >= 0.5:
                row.append(info)
                placed = True
                break
        if not placed:
            rows.append([info])
    ordered: list[ShapeInfo] = []
    for row in rows:
        ordered.extend(sorted(row, key=lambda i: (i.left, i.top)))
    for index, info in enumerate(ordered):
        info.reading_order = index
    return ordered


def _has_text(shape) -> bool:
    return bool(getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip())


def _shape_text(shape) -> str:
    return clean(shape.text_frame.text.replace("\n", " "))


def _max_font_pt(shape) -> float:
    """텍스트 상자의 가장 큰 명시적 글꼴 크기(pt). 상속만 있으면 0."""
    best = 0.0
    try:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if run.font.size is not None:
                    best = max(best, run.font.size.pt)
            if para.font.size is not None:
                best = max(best, para.font.size.pt)
    except Exception:
        pass
    return best


def _body_font_pt(infos: list[ShapeInfo]) -> float:
    """슬라이드 본문 글씨 크기의 추정값: 글이 있는 상자들의 명시적 글꼴 크기 중앙값. 명시된 크기가 없으면 0."""
    sizes = []
    for info in infos:
        if _has_text(info.shape):
            pt = _max_font_pt(info.shape)
            if pt > 0:
                sizes.append(pt)
    sizes.sort()
    return sizes[len(sizes) // 2] if sizes else 0.0


def _title_by_font(infos: list[ShapeInfo], slide_height: int) -> ShapeInfo | None:
    """위쪽 영역의 텍스트 상자 중 글씨가 가장 큰 것. 크기 정보가 없으면 가장 위의 것.
    본문 글씨의 0.8배 미만인 작은 상자(문서명·대외비 같은 머리글)는 건너뛴다 (PDF 리더의 first-line 과 같은 기준)."""
    body = _body_font_pt(infos)
    cands = []
    for info in infos:
        sh = info.shape
        if not _has_text(sh):
            continue
        text = _shape_text(sh)
        if not text or len(text) > TITLE_MAX_CHARS:
            continue
        if slide_height and info.top > slide_height * TITLE_ZONE:
            continue
        pt = _max_font_pt(sh)
        if 0 < pt < TITLE_MIN_RATIO * body:
            continue
        cands.append((-pt, info.top, info.reading_order, info))
    if not cands:
        return None
    cands.sort(key=lambda t: t[:3])
    return cands[0][3]


def _graphic_uri(shape) -> str | None:
    data = shape._element.find(".//" + A + "graphicData")
    return None if data is None else data.get("uri")


def _paragraph_text(p) -> str:
    """a:p 의 텍스트. 줄바꿈(a:br)은 공백으로."""
    parts: list[str] = []
    for child in p:
        if child.tag in (A + "r", A + "fld"):
            node = child.find(A + "t")
            if node is not None and node.text:
                parts.append(node.text)
        elif child.tag == A + "br":
            parts.append(" ")
    return "".join(parts)


def _smartart_lines(slide, shape) -> list[Line]:
    """SmartArt 데이터 모델(diagram data part)의 점(pt)마다 글을 꺼낸다. 전환·배치용 점은 건너뛴다."""
    from lxml import etree

    rel_ids = shape._element.find(".//" + DGM + "relIds")
    rid = rel_ids.get(R + "dm") if rel_ids is not None else None
    if not rid:
        return []
    try:
        part = slide.part.related_part(rid)
        root = etree.fromstring(part.blob, etree.XMLParser(resolve_entities=False, no_network=True))
    except Exception:
        return []
    out: list[Line] = []
    for point in root.iter(DGM + "pt"):
        if point.get("type") in ("pres", "parTrans", "sibTrans"):
            continue
        body = point.find(DGM + "t")
        if body is None:
            continue
        for p in body.findall(A + "p"):
            text = clean(_paragraph_text(p))
            if text:
                out.append(Line(idx=0, text=text, source="diagram"))
    return out


def layout_report(path: Path, pages: set[int] | None = None) -> list[str]:
    """진단용: 슬라이드마다 고른 제목과 도형 목록(읽기 순서, 종류, 위치 %, 글꼴 크기, 글자 수, 앞부분)."""
    prs = Presentation(str(path))
    w, h = float(prs.slide_width or 1), float(prs.slide_height or 1)
    doc = read_pptx(path)
    out = [f"== {path}", *[f"  안내: {n}" for n in doc.notes]]
    for sno, slide in enumerate(prs.slides, 1):
        if pages and sno not in pages:
            continue
        pg = doc.pages[sno - 1]
        out.append(f"-- {sno}쪽  제목: {pg.title[:60] or '(없음)'}  [{pg.title_source}]")
        for info in assign_reading_order(collect_shapes(slide.shapes)):
            sh = info.shape
            kind = ("표" if getattr(sh, "has_table", False) and sh.has_table else "차트" if getattr(sh, "has_chart", False) and sh.has_chart
                    else "SmartArt" if _graphic_uri(sh) == DIAGRAM_URI else "그림" if sh._element.tag == P + "pic"
                    else "글상자" if _has_text(sh) else str(sh.shape_type or "도형"))
            text = _shape_text(sh) if _has_text(sh) else ""
            out.append(f"     {info.reading_order:2d} {kind:<8} top {100 * info.top / h:4.0f}% left {100 * info.left / w:4.0f}%"
                       f"  {_max_font_pt(sh) or 0:4.0f}pt{' 그룹' if info.in_group else '    '}  {len(text):3d}자  {text[:50]}")
    return out


def read_pptx(path: Path) -> Document:
    prs = Presentation(str(path))
    slide_height = prs.slide_height or 0
    pages: list[Page] = []
    hidden: list[int] = []
    smartart = charts = 0
    for sno, slide in enumerate(prs.slides, 1):
        if slide._element.get("show") in ("0", "false"):
            hidden.append(sno)
        title, title_id, title_source = "", None, "없음"
        ts = slide.shapes.title
        if ts is not None and ts.has_text_frame:
            title = _shape_text(ts)
            title_id = ts.shape_id
            title_source = "placeholder"
        infos = assign_reading_order(collect_shapes(slide.shapes))
        if not title:
            cand = _title_by_font(infos, slide_height)
            if cand is not None:
                title, title_id, title_source = _shape_text(cand.shape), cand.shape.shape_id, "font"
        lines: list[Line] = []
        has_graphics = False
        for info in infos:
            sh = info.shape
            if title_id is not None and sh.shape_id == title_id:
                continue
            uri = _graphic_uri(sh)
            if getattr(sh, "has_table", False) and sh.has_table:
                lines.extend(table_lines([[c.text for c in row.cells] for row in sh.table.rows]))
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                has_graphics = True
                charts += 1
                try:
                    chart = sh.chart
                    if chart.has_title and chart.chart_title.has_text_frame:
                        text = clean(chart.chart_title.text_frame.text.replace("\n", " "))
                        if text:
                            lines.append(Line(idx=0, text=text, source="diagram"))
                except Exception:
                    pass
            elif uri == DIAGRAM_URI:
                has_graphics = True
                smartart += 1
                lines.extend(_smartart_lines(slide, sh))
            elif uri == OLE_URI or sh._element.tag == P + "pic":
                has_graphics = True
            elif _has_text(sh):
                for para in sh.text_frame.paragraphs:
                    text = clean("".join(r.text for r in para.runs) or para.text)
                    if text:
                        lines.append(Line(idx=0, text=text, level=para.level or 0))
                if sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM) and info.in_group:
                    has_graphics = True           # 블록도 안의 글상자
            elif sh.shape_type in _GRAPHIC_TYPES:
                has_graphics = True
        if not title and lines:
            title, lines, title_source = lines[0].text, lines[1:], "first-line"
        pages.append(Page(no=sno, title=title, lines=lines, has_graphics=has_graphics, title_source=title_source))
    notes: list[str] = []
    if hidden:
        notes.append(f"숨긴 슬라이드 {len(hidden)}장도 읽음 ({', '.join(str(n) for n in hidden[:8])}{' …' if len(hidden) > 8 else ''})")
    if smartart:
        notes.append(f"SmartArt {smartart}개의 글을 도형 줄로 읽음")
    if charts:
        notes.append(f"차트 {charts}개는 제목만 읽음 (데이터·범례 제외)")
    return Document(path=Path(path), kind="pptx", pages=pages, title=pages[0].title if pages else "", notes=notes)

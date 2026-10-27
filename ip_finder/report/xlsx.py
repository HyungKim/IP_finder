"""Excel 리포트: 검토 목록 / 보류 / 2차 제외 / 페이지 / 변경 사항 / 제외 페이지 / 이미지 페이지 / 1차 제외 (참고) / 실행 정보.
'2차 제외' = 분류기 1차에서 뽑혔지만 LLM 2차가 제외한 문장. 잘못 제외된 것을 피드백 열에 A~D 로 적어 살린다.
문장 시트마다 '1차 분류기'와 '2차 LLM' 열이 있어 어느 단계가 무엇으로 판정했는지 바로 보인다.
검토 목록·보류·제외 시트에는 '피드백' 열이 있어 O / X / A~D 를 적으면 ip-finder feedback 이 골드 시트로 가져간다."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..models import Document, LabeledLine, RunInfo
from ..post import in_hold_list, in_review_list, is_stage1_out, is_stage2_out
from ..post.templates import discussion_point
from ..prep.page_rules import IMAGE_PAGE

FILLS = {
    "A": PatternFill("solid", fgColor="FFE0B2"),
    "B": PatternFill("solid", fgColor="FFF9C4"),
    "C": PatternFill("solid", fgColor="DCEDC8"),
    "D": PatternFill("solid", fgColor="BBDEFB"),
    "X?": PatternFill("solid", fgColor="EEEEEE"),
    "?": PatternFill("solid", fgColor="FFCDD2"),
}
HEAD_FILL = PatternFill("solid", fgColor="37474F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
FEEDBACK_FILL = PatternFill("solid", fgColor="FFFDE7")
WRAP = Alignment(wrap_text=True, vertical="top")
SOURCE_NAME = {"clf": "분류기", "llm": "LLM", "verified": "분류기+LLM", "gold": "골드"}
FEEDBACK_HELP = "O=맞음, X=틀림(제외), A/B/C/D=올바른 등급"

SHEET_REVIEW, SHEET_HOLD, SHEET_LLM_OUT, SHEET_CLF_OUT = "검토 목록", "보류", "2차 제외", "1차 제외 (참고)"
LINE_HEADERS = ["번호", "섹션", "페이지", "줄", "등급", "발췌", "논의 포인트", "점수", "가중치", "1차 분류기", "2차 LLM", "판정", "비고", "피드백", "메모"]
LINE_WIDTHS = [6, 14, 8, 5, 6, 60, 34, 7, 7, 9, 8, 9, 28, 8, 24]
FEEDBACK_COL = LINE_HEADERS.index("피드백") + 1


def stage_labels(ll: LabeledLine) -> tuple[str, str]:
    """(1차 분류기 라벨, 2차 LLM 라벨). 해당 단계를 거치지 않았으면 '-'."""
    return (ll.clf_label or "-", ll.llm_label or "-")


def is_llm_rejected(ll: LabeledLine) -> bool:
    """1차를 통과했지만 2차(LLM, 또는 보정·골드)가 대상 등급이 아니라고 본 문장. '2차 제외' 시트."""
    return is_stage2_out(ll)


def is_clf_rejected(ll: LabeledLine) -> bool:
    """분류기 1차에서 대상 등급이 아니라고 본 문장 (X? 는 보류 시트). '1차 제외 (참고)' 시트."""
    return is_stage1_out(ll)


def is_held(ll: LabeledLine) -> bool:
    """보류 시트: 점수 미달, 1차 단계의 X?, 판정 실패."""
    return in_hold_list(ll)


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.freeze_panes = "A2"
    return ws


def _finish(ws, n_cols: int, feedback_col: int | None = None) -> None:
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{ws.max_row}"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    if feedback_col and ws.max_row > 1:
        col = get_column_letter(feedback_col)
        dv = DataValidation(type="list", formula1='"O,X,A,B,C,D"', allow_blank=True, showErrorMessage=False)
        dv.promptTitle, dv.prompt = "피드백", FEEDBACK_HELP
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ws.max_row}")
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=feedback_col).fill = FEEDBACK_FILL
            ws.cell(row=r, column=feedback_col + 1).fill = FEEDBACK_FILL


def _line_row(n: int, ll: LabeledLine, extra: list | None = None) -> list:
    first, second = stage_labels(ll)
    return [n, ll.section, ll.page_no, ll.line_idx, ll.label, ll.text, discussion_point(ll.label), ll.score, ll.page_weight,
            first, second, SOURCE_NAME.get(ll.source, ll.source), ", ".join(ll.notes), "", ""] + (extra or [])


def _line_sheet(wb, title, rows: list[LabeledLine], first=False, extra_header: str | None = None, extra_of=None, note: str | None = None):
    headers = LINE_HEADERS + ([extra_header] if extra_header else [])
    widths = LINE_WIDTHS + ([40] if extra_header else [])
    ws = _sheet(wb, title, headers, widths, first=first)
    for n, ll in enumerate(rows, 1):
        ws.append(_line_row(n, ll, [extra_of(ll)] if extra_of else None))
        ws.cell(row=ws.max_row, column=5).fill = FILLS.get(ll.label, FILLS["X?"]) if ll.label in FILLS else PatternFill()
    if not rows and note:
        ws.append(["", "", "", "", "", note])
    _finish(ws, len(headers), FEEDBACK_COL)
    return ws


def write_report(out: Path, doc: Document, labeled: list[LabeledLine], info: RunInfo,
                 removed: Counter | None = None, changes: list[dict] | None = None) -> Path:
    wb = Workbook()
    cascade = info.mode == "cascade"

    _line_sheet(wb, SHEET_REVIEW, [ll for ll in labeled if in_review_list(ll)], first=True,
                note="검토 목록에 오른 문장이 없습니다.")

    def hold_reason(ll: LabeledLine) -> str:
        return ll.hold_reason or ("모델은 제외, 기술 단서어 있음" if ll.label == "X?" else "판정 실패")

    _line_sheet(wb, SHEET_HOLD, [ll for ll in labeled if is_held(ll)], extra_header="보류 사유", extra_of=hold_reason,
                note="보류된 문장이 없습니다.")

    _line_sheet(wb, SHEET_LLM_OUT, [ll for ll in labeled if is_llm_rejected(ll)],
                note=(f"분류기 1차에서 뽑힌 문장 중 LLM 2차가 {info.target_grades} 가 아니라고 본 문장이 없습니다." if cascade else
                      "이 모드에서는 LLM 2차 판정이 없습니다 (cascade 모드에서 채워짐)." if info.mode == "fast" else "LLM 이 제외한 문장이 없습니다."))

    ws = _sheet(wb, "페이지", ["페이지", "제목", "관련성", "가중치", "근거", "상태", "문장", "목록", "보류", "1차 제외", "2차 제외", "제목 출처"],
                [8, 48, 8, 8, 22, 26, 7, 7, 7, 8, 8, 10])
    by_page: dict[int, list[LabeledLine]] = {}
    for ll in labeled:
        by_page.setdefault(ll.page_no, []).append(ll)
    for p in doc.pages:
        rows = by_page.get(p.no, [])
        ws.append([p.no, p.title, p.relevance, p.weight, p.weight_reason, p.excluded_reason or "판정", len(rows),
                   sum(1 for ll in rows if in_review_list(ll)), sum(1 for ll in rows if is_held(ll)),
                   sum(1 for ll in rows if is_clf_rejected(ll)), sum(1 for ll in rows if is_llm_rejected(ll)), p.title_source])
    _finish(ws, 12)

    ws = _sheet(wb, "변경 사항", ["변경", "페이지", "줄", "발췌", "이전", "현재", "목록", "판정", "설명"], [10, 8, 5, 64, 8, 8, 6, 11, 44])
    if info.prev_run_id:
        ws.append(["기준", "", "", f"직전 실행 {info.prev_run_id} 과 비교 (현재 {info.run_id})", "", "", "", "", ""])
        for s in info.settings_diff or ["설정 변경 없음"]:
            ws.append(["설정", "", "", s, "", "", "", "", ""])
        for r in changes or []:
            ws.append([r["kind"], r["page"], r["line"], r["text"], r["old"], r["new"], r["in_list"], r["source"], r["desc"]])
        if not changes:
            ws.append(["", "", "", "변경된 문장 없음", "", "", "", "", ""])
    else:
        ws.append(["", "", "", "이 파일의 이전 실행 기록이 없어 비교하지 않았습니다. 다음 실행부터 여기에 달라진 행이 적힙니다.", "", "", "", "", ""])
    _finish(ws, 9)

    ws = _sheet(wb, "제외 페이지", ["페이지", "제목", "사유"], [8, 50, 30])
    for p in doc.pages:
        if p.excluded_reason and p.excluded_reason != IMAGE_PAGE:
            ws.append([p.no, p.title, p.excluded_reason])
    _finish(ws, 3)

    ws = _sheet(wb, "이미지 페이지", ["페이지", "제목", "비고"], [8, 50, 40])
    for p in doc.pages:
        if p.excluded_reason == IMAGE_PAGE:
            ws.append([p.no, p.title, "텍스트 없음. 도면·그림이면 직접 확인"])
    _finish(ws, 3)

    _line_sheet(wb, SHEET_CLF_OUT, [ll for ll in labeled if is_clf_rejected(ll)],
                note=("분류기 1차에서 제외된 문장이 없습니다." if info.mode != "llm" else "llm 모드에서는 분류기 1차가 없습니다."))

    ws = _sheet(wb, "실행 정보", ["항목", "값"], [24, 90])
    grade_counts = Counter(ll.label for ll in labeled)
    list_counts = Counter(ll.label for ll in labeled if in_review_list(ll))
    n_clf_out = sum(1 for ll in labeled if is_clf_rejected(ll))
    n_llm_out = sum(1 for ll in labeled if is_llm_rejected(ll))
    n_cand = sum(1 for ll in labeled if ll.llm_label)
    rows = [
        ("입력 파일", info.source), ("실행 ID", info.run_id or "-"), ("직전 실행 ID", info.prev_run_id or "-"),
        ("판정 모드", info.mode),
        ("추출 등급", info.target_grades),
        ("판정 흐름", {"cascade": f"분류기 1차: {info.target_grades} 로 판정된 문장만 후보 → LLM 2차: 후보가 {info.target_grades} 인지 판정해 최종 등급 → 두 단계 통과만 검토 목록",
                   "fast": f"분류기만 ({info.target_grades} 만 목록)", "llm": f"전부 LLM ({info.target_grades} 만 목록)"}.get(info.mode, info.mode)),
        ("단계별 문장 수", f"전체 {len(labeled)} → 1차 제외 {n_clf_out} → 2차 후보 {n_cand} → 2차 제외 {n_llm_out} → 보류 {sum(1 for ll in labeled if is_held(ll))} → 검토 목록 {sum(list_counts.values())}"),
        ("LLM 모델", info.model), ("프롬프트 버전", info.prompt_version),
        ("임베딩 모델", info.embed_model or "-"), ("분류기", info.classifier or "-"),
        ("페이지 가중치", f"{'켬' if info.page_weight_enabled else '끔'} (관련성 낮음 {info.low_relevance_weight}, 낮음 페이지 {info.pages_low}개)"),
        ("목록 점수 기준", info.accept_score), ("점수 미달 보류", info.held_lines),
        ("골드 정답 반영", f"{info.gold_lines}문장 (모델 판정과 달랐던 것 {info.gold_changed}) / 읽은 골드: {info.gold_desc}" if info.gold_desc else "없음 (골드 파일이 없거나 use_gold=false)"),
        ("임베딩 시간(초)", round(info.embed_sec, 1)), ("LLM 시간(초)", round(info.llm_sec, 1)),
        ("시작", info.started_at), ("소요 시간(초)", round(info.elapsed_sec, 1)),
        ("페이지 수", info.pages_total), ("라벨링 페이지", info.pages_labeled),
        ("제외 페이지", info.pages_excluded), ("이미지 페이지", info.pages_image),
        ("LLM 호출", info.llm_calls), ("캐시 적중", info.cache_hits),
        ("검토 목록", ", ".join(f"{g}: {list_counts.get(g, 0)}" for g in info.target_grades.split(",")) + f" (합 {sum(list_counts.values())})"),
        ("등급 분포(전체)", ", ".join(f"{g}: {grade_counts.get(g, 0)}" for g in ("A", "B", "C", "D", "X", "X?", "?"))),
        ("직전 실행 대비", ", ".join(f"{k}: {v}" for k, v in (info.changes or {}).items() if k != "total" and v) or ("변경 없음" if info.prev_run_id else "-")),
        ("달라진 설정", "; ".join(info.settings_diff) or "-"),
        ("전처리로 제거된 줄", ", ".join(f"{k}: {v}" for k, v in (removed or {}).items()) or "없음"),
        ("경고", "; ".join(info.warnings) or "없음"),
        ("피드백 방법", f"검토 목록·보류·제외 시트의 '피드백' 열에 {FEEDBACK_HELP}. 저장 후 UI '정확도 관리 → 피드백 가져오기' 또는 ip-finder feedback <이 파일>"),
    ]
    for r in rows:
        ws.append(list(r))
    _finish(ws, 2)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return out

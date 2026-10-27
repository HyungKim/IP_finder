"""결과 Excel의 '피드백' 열을 골드 시트로 가져온다.

피드백 값: O(맞음) / X(틀림, 제외) / A·B·C·D(올바른 등급). 비어 있으면 무시 (confirm_unmarked=True 면 검토 목록의 빈 행을 '맞음'으로).
골드 파일이 이미 있으면 같은 (파일, 페이지, 줄, 문장) 행은 정답을 덮어쓰고, 없는 행은 뒤에 붙인다."""
from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .. import PROMPT_VERSION
from ..history import History
from ..goldmap import GOLD_HEADERS, GOLD_SUFFIX, load_gold
from .gold import write_gold_workbook

_OK = {"O", "o", "0", "ㅇ", "맞음", "✓", "V", "v", "Y", "y"}
_NO = {"X", "x", "ㅌ", "틀림", "제외", "N", "n"}
FEEDBACK_SHEETS = ("검토 목록", "보류", "2차 제외", "1차 제외 (참고)", "제외 문장")   # 마지막은 v0.4 결과 파일 호환


def _read_sheet(ws) -> list[dict]:
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(h) for h in next(rows_iter)]
    except StopIteration:
        return []
    idx = {h: i for i, h in enumerate(header)}
    out = []
    for r in rows_iter:
        if r is None or idx.get("발췌") is None or r[idx["발췌"]] in (None, ""):
            continue
        out.append({h: r[i] for h, i in idx.items()})
    return out


def interpret(feedback, current_label: str) -> str | None:
    """피드백 값을 정답 등급으로. 해석할 수 없으면 None."""
    s = str(feedback or "").strip()
    if not s:
        return None
    up = s.upper()
    if up in ("A", "B", "C", "D"):
        return up
    if s in _OK or up in _OK:
        return "X" if current_label in ("X", "X?", "?") else current_label
    if s in _NO or up in _NO:
        return "X"
    return None


def read_feedback(result_xlsx: Path, confirm_unmarked: bool = False) -> tuple[list[dict], dict]:
    """결과 Excel에서 피드백이 있는 행을 읽어 골드 행(dict)으로 만든다."""
    wb = load_workbook(str(result_xlsx), read_only=True, data_only=True)
    titles: dict[int, str] = {}
    if "페이지" in wb.sheetnames:
        for r in _read_sheet_simple(wb["페이지"]):
            try:
                titles[int(r.get("페이지"))] = str(r.get("제목") or "")
            except (TypeError, ValueError):
                continue
    info = {}
    if "실행 정보" in wb.sheetnames:
        for r in wb["실행 정보"].iter_rows(min_row=2, values_only=True):
            if r and r[0]:
                info[str(r[0])] = r[1]
    source_name = Path(str(info.get("입력 파일") or result_xlsx.stem.replace("_검토후보", ""))).name
    rows, skipped, corrections = [], 0, Counter()
    for sheet in FEEDBACK_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        for r in _read_sheet(wb[sheet]):
            label = str(r.get("등급") or "").strip()
            fb = r.get("피드백")
            gold = interpret(fb, label)
            if gold is None:
                if str(fb or "").strip():
                    skipped += 1
                    continue
                if confirm_unmarked and sheet == "검토 목록" and label in ("A", "B", "C", "D"):
                    gold = label
                else:
                    continue
            try:
                page_no = int(r.get("페이지"))
                line_idx = int(r.get("줄") or 0)
            except (TypeError, ValueError):
                skipped += 1
                continue
            norm_label = "X" if label in ("X?", "?") else label
            corrections[f"{norm_label}→{gold}" if norm_label != gold else f"{gold} 확인"] += 1
            rows.append({"파일": source_name, "페이지": page_no, "제목": titles.get(page_no, ""), "줄": line_idx,
                         "문장": str(r.get("발췌")), "모델 라벨": norm_label, "후처리 라벨": label, "판정": r.get("판정") or "",
                         "확신": "", "점수": r.get("점수") if r.get("점수") not in (None, "") else "", "가중치": r.get("가중치") or "",
                         "정답": gold, "메모": str(r.get("메모") or "")})
    return rows, {"source_name": source_name, "skipped": skipped, "corrections": dict(corrections),
                  "run_id": str(info.get("실행 ID") or ""), "model": str(info.get("LLM 모델") or "")}


def _read_sheet_simple(ws) -> list[dict]:
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(h) for h in next(rows_iter)]
    except StopIteration:
        return []
    return [dict(zip(header, r)) for r in rows_iter if r and r[0] is not None]


def default_gold_path(result_xlsx: Path, source_name: str, gold_dir: Path | None = None) -> Path:
    """기본: 설정의 골드 폴더(data/gold)/<보고서 이름>_gold.xlsx. gold_dir 이 없으면 결과 파일 옆."""
    stem = Path(source_name).stem if source_name else result_xlsx.stem.replace("_검토후보", "")
    return (gold_dir or result_xlsx.parent) / f"{stem}{GOLD_SUFFIX}"


def import_feedback(result_xlsx: Path, gold_out: Path | None = None, confirm_unmarked: bool = False,
                    history: History | None = None, gold_dir: Path | None = None) -> dict:
    new_rows, meta = read_feedback(result_xlsx, confirm_unmarked)
    gold_out = gold_out or default_gold_path(result_xlsx, meta["source_name"], gold_dir)
    existing: list[dict] = load_gold(gold_out) if gold_out.exists() else []
    by_key = {(str(r.get("파일")), int(r.get("페이지") or 0), int(r.get("줄") or 0), str(r.get("문장"))): r for r in existing}
    updated, added = 0, 0
    for r in new_rows:
        k = (r["파일"], r["페이지"], r["줄"], r["문장"])
        if k in by_key:
            by_key[k]["정답"] = r["정답"]
            if r["메모"]:
                by_key[k]["메모"] = r["메모"]
            updated += 1
        else:
            by_key[k] = r
            added += 1
    rows = []
    for n, r in enumerate(by_key.values(), 1):
        rows.append([n] + [r.get(h, "") if r.get(h, "") is not None else "" for h in GOLD_HEADERS[1:]])
    meta_rows = [("파일", meta["source_name"]), ("피드백 원본", str(result_xlsx)), ("프롬프트 버전", PROMPT_VERSION),
                 ("갱신", datetime.now().isoformat(timespec="seconds")), ("행 수", str(len(rows)))]
    write_gold_workbook(gold_out, rows, meta_rows)
    record = {"result": str(result_xlsx), "gold": str(gold_out), "source_name": meta["source_name"], "run_id": meta["run_id"],
              "feedback_rows": len(new_rows), "added": added, "updated": updated, "skipped": meta["skipped"],
              "corrections": meta["corrections"], "gold_rows": len(rows)}
    (history or History(None)).log_feedback(record)
    return record

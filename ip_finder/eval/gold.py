"""골드 세트 라벨링 시트. 모델 초안 라벨을 '정답' 열에 미리 채워 두고, 사용자는 틀린 것만 고친다.
이 파일은 실제 보고서 문장을 담으므로 회사 PC의 프로젝트 바깥 폴더에만 둔다."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .. import PROMPT_VERSION
from ..config import Config
from ..pipeline import run_file
from .score import GOLD_HEADERS

RUBRIC_PATH = Path(__file__).parent.parent / "label" / "prompts" / "rubric.md"
GOLD_WIDTHS = (6, 18, 8, 24, 5, 70, 10, 10, 11, 8, 7, 7, 8, 30)
SOURCE_NAME = {"clf": "분류기", "llm": "LLM", "verified": "분류기+LLM"}

GUIDE = [
    "정답 열에 A / B / C / D / X 중 하나를 적습니다. 모델 라벨이 미리 채워져 있으니 틀린 것만 고치세요.",
    "X는 기술 내용이 없거나 특허팀과 논의할 가치가 없는 문장입니다. 정밀도 우선이므로 판정이 어려우면 X로 두고 메모에 이유를 적으세요.",
    "채점: ip-finder eval <이 파일>   (현재 설정으로 다시 판정해 채점)   /   --stored (저장된 라벨을 채점)",
    "학습: ip-finder train <이 파일> [다른 골드 파일 ...]   → 분류기 저장 후 hybrid/fast 모드가 자동으로 사용",
    "결과 Excel의 '피드백' 열에 적은 내용은 ip-finder feedback <결과.xlsx> 로 이 형식의 골드 시트에 합쳐집니다.",
]


def write_gold_workbook(out: Path, rows: list[list], meta: list[tuple[str, str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "라벨링"
    ws.append(GOLD_HEADERS)
    for c in range(1, len(GOLD_HEADERS) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="CFD8DC")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(GOLD_WIDTHS, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("안내")
    for line in [*GUIDE, "", *RUBRIC_PATH.read_text(encoding="utf-8").splitlines()]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 120

    ws3 = wb.create_sheet("메타")
    for k, v in meta:
        ws3.append([k, v])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return out


def make_gold_sheet(path: Path, cfg: Config, out: Path | None = None, client=None, embedder=None) -> Path:
    res = run_file(path, cfg, client, embedder=embedder)
    out = out or path.with_name(f"{path.stem}_gold.xlsx")
    rows = []
    for n, ll in enumerate(res.labeled, 1):
        answer = "X" if ll.raw_label in ("X", "X?") else ("" if ll.raw_label == "?" else ll.raw_label)
        rows.append([n, path.name, ll.page_no, ll.page_title, ll.line_idx, ll.text, ll.raw_label, ll.label,
                     SOURCE_NAME.get(ll.source, ll.source), round(ll.clf_conf, 3) if ll.clf_conf is not None else "",
                     float(ll.score), float(ll.page_weight), answer, ""])
    meta = [("파일", str(path)), ("판정 모드", res.info.mode), ("LLM 모델", cfg.model), ("임베딩 모델", res.info.embed_model or "-"),
            ("분류기", res.info.classifier or "-"), ("프롬프트 버전", PROMPT_VERSION), ("생성", res.info.started_at)]
    return write_gold_workbook(out, rows, meta)

"""골드 시트 채점: 정확도, 등급별 정밀도·재현율·F1, 혼동 행렬. --stored 가 아니면 현재 모델·프롬프트로 다시 라벨링한다."""
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from ..cache import LabelCache
from ..config import Config
from ..models import GRADES, Line, Page, RunInfo
from ..post import apply_safety_net, apply_title_fix

GOLD_HEADERS = ["번호", "파일", "페이지", "제목", "줄", "문장", "모델 라벨", "후처리 라벨", "판정", "확신", "정답", "메모"]

LABELS = list(GRADES)


def _norm(label) -> str:
    s = str(label or "").strip().upper()
    return "X" if s in ("X?", "X") else s


def load_gold(path: Path) -> list[dict]:
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb["라벨링"]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows_iter)]
    idx = {h: i for i, h in enumerate(header)}
    rows = []
    for r in rows_iter:
        if r is None or r[idx["문장"]] is None:
            continue
        gold = _norm(r[idx["정답"]])
        if gold not in LABELS:
            continue
        rows.append({h: r[idx[h]] for h in GOLD_HEADERS if h in idx} | {"정답": gold})
    return rows


def score(pairs: list[tuple[str, str]]) -> dict:
    """pairs = [(정답, 예측)]"""
    total = len(pairs)
    correct = sum(1 for g, p in pairs if g == p)
    conf: dict[str, Counter] = defaultdict(Counter)
    for g, p in pairs:
        conf[g][p] += 1
    per = {}
    for lab in LABELS:
        tp = conf[lab][lab]
        fn = sum(conf[lab].values()) - tp
        fp = sum(conf[g][lab] for g in LABELS if g != lab)
        prec = tp / (tp + fp) if tp + fp else 0.0
        rec = tp / (tp + fn) if tp + fn else 0.0
        f1 = 2 * prec * rec / (prec + rec) if prec + rec else 0.0
        per[lab] = {"support": tp + fn, "precision": prec, "recall": rec, "f1": f1}
    return {"total": total, "correct": correct, "accuracy": correct / total if total else 0.0,
            "per_grade": per, "confusion": {g: dict(conf[g]) for g in LABELS}}


def evaluate(gold_path: Path, cfg: Config, rerun: bool = True, client=None, embedder=None, classifier=None) -> dict:
    rows = load_gold(gold_path)
    if not rows:
        raise ValueError("채점할 행이 없습니다. '정답' 열을 확인하세요.")
    info = RunInfo(source=str(gold_path), model=cfg.model, prompt_version="", started_at="")
    if rerun:
        from ..pipeline import Services, label_pages
        groups: dict[tuple, list[dict]] = defaultdict(list)
        for r in rows:
            groups[(r.get("파일"), r.get("페이지"))].append(r)
        pages, order = [], []
        for n, ((_, pno), grp) in enumerate(groups.items(), 1):
            grp.sort(key=lambda r: int(r.get("줄") or 0))
            lines = [Line(idx=i, text=str(r["문장"])) for i, r in enumerate(grp, 1)]
            pages.append(Page(no=n, title=str(grp[0].get("제목") or ""), lines=lines))
            order.extend(grp)
        cache = LabelCache(cfg.cache_file() if cfg.use_cache else None)
        try:
            svc = Services(cfg, cache, client=client, embedder=embedder, classifier=classifier)
            labeled = label_pages(pages, cfg, svc, info)
        finally:
            cache.close()
        if cfg.title_fix:
            apply_title_fix(labeled)
        if cfg.safety_net:
            apply_safety_net(labeled)
        raw_pairs = [(r["정답"], _norm(ll.raw_label)) for r, ll in zip(order, labeled)]
        final_pairs = [(r["정답"], _norm(ll.label)) for r, ll in zip(order, labeled)]
        source = f"재실행: 모드 {info.mode}, LLM {cfg.model}" + (f", 임베딩 {info.embed_model}, {info.classifier}" if info.embed_model else "")
    else:
        raw_pairs = [(r["정답"], _norm(r.get("모델 라벨"))) for r in rows]
        final_pairs = [(r["정답"], _norm(r.get("후처리 라벨"))) for r in rows]
        source = "시트에 저장된 라벨"
    return {"source": source, "rows": len(rows), "raw": score(raw_pairs), "final": score(final_pairs),
            "llm_lines": info.llm_lines, "clf_lines": info.clf_lines, "llm_calls": info.llm_calls}


def print_report(rep: dict) -> None:
    print(f"골드 {rep['rows']}문장 / {rep['source']}")
    if rep.get("llm_lines") or rep.get("clf_lines"):
        print(f"  분류기 채택 {rep.get('clf_lines', 0)}문장, LLM 판정 {rep.get('llm_lines', 0)}문장, LLM 호출 {rep.get('llm_calls', 0)}회")
    for name, s in (("모델 라벨", rep["raw"]), ("후처리 후", rep["final"])):
        print(f"\n[{name}] 정확도 {s['correct']}/{s['total']} = {s['accuracy']:.1%}")
        print("  등급  건수  정밀도  재현율   F1")
        for lab in LABELS:
            m = s["per_grade"][lab]
            print(f"  {lab:<4} {m['support']:>4}  {m['precision']:6.1%} {m['recall']:6.1%} {m['f1']:6.1%}")
        print("  혼동 행렬 (행=정답, 열=예측): " + " ".join(f"{l:>3}" for l in LABELS))
        for g in LABELS:
            print(f"    {g:<3} " + " ".join(f"{s['confusion'][g].get(p, 0):>3}" for p in LABELS))

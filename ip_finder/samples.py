"""설명용 가짜 보고서 생성기. demo/sample.json → PPTX, PDF. 실제 사내 자료는 포함하지 않는다.
사용: python -m ip_finder.samples [출력폴더]   (기본 demo/)"""
from __future__ import annotations

import json
import sys
from pathlib import Path

FOOTER = "PS-200 개선 보고 | 대외비 | 품질혁신팀"
TABLE_PAGE = 8
TABLE = [["항목", "개선 전", "개선 후"], ["오프셋 드리프트", "0.35 %FS", "0.10 %FS"], ["조립 수율", "92.1 %", "97.8 %"]]


def load_sample(path: Path | None = None) -> dict:
    path = path or Path(__file__).resolve().parent.parent / "demo" / "sample.json"
    return json.loads(path.read_text(encoding="utf-8"))


def make_pptx(sample: dict, out: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    for pg in sample["pages"]:
        if pg["no"] == 1:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = pg["title"]
            slide.placeholders[1].text = "\n".join(pg["lines"])
        elif not pg["lines"]:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = pg["title"]
            slide.shapes.add_shape(1, Inches(1), Inches(2), Inches(3.5), Inches(3))   # 도면 대신 사각형
            slide.shapes.add_shape(1, Inches(5.5), Inches(2), Inches(3.5), Inches(3))
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = pg["title"]
            tf = slide.placeholders[1].text_frame
            tf.text = pg["lines"][0]
            for line in pg["lines"][1:]:
                p = tf.add_paragraph()
                p.text = line
            if pg["no"] == TABLE_PAGE:
                rows, cols = len(TABLE), len(TABLE[0])
                table = slide.shapes.add_table(rows, cols, Inches(1), Inches(5.2), Inches(8), Inches(1.2)).table
                for r in range(rows):
                    for c in range(cols):
                        table.cell(r, c).text = TABLE[r][c]
        for note in pg.get("notes", []):
            slide.notes_slide.notes_text_frame.text = note
        box = slide.shapes.add_textbox(Inches(0.3), Inches(7.0), Inches(6), Inches(0.3))
        box.text_frame.text = FOOTER
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def _font():
    import pymupdf
    for name in ("cjk",):
        try:
            return pymupdf.Font(name)
        except Exception:
            pass
    for path in ("/System/Library/Fonts/Supplemental/AppleGothic.ttf", "/System/Library/Fonts/AppleSDGothicNeo.ttc",
                 "C:/Windows/Fonts/malgun.ttf", "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
                 "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"):
        if Path(path).exists():
            return pymupdf.Font(fontfile=path)
    raise RuntimeError("한글 폰트를 찾지 못했습니다")


def make_pdf(sample: dict, out: Path) -> Path:
    import pymupdf

    font = _font()
    doc = pymupdf.open()
    for pg in sample["pages"]:
        page = doc.new_page(width=842, height=595)   # 가로 A4
        tw = pymupdf.TextWriter(page.rect)
        tw.append((50, 70), pg["title"], font=font, fontsize=22)
        y = 130
        for line in pg["lines"]:
            tw.append((70, y), f"• {line}", font=font, fontsize=13)
            y += 34
        if pg["no"] == TABLE_PAGE:
            for r in TABLE:
                tw.append((70, y), "   ".join(f"{c:<14}" for c in r), font=font, fontsize=12)
                y += 26
        tw.append((50, 570), FOOTER, font=font, fontsize=8)
        tw.write_text(page)
        if not pg["lines"]:
            for x in (100, 460):
                page.draw_rect(pymupdf.Rect(x, 150, x + 280, 420), color=(0, 0, 0), width=1.5)
                page.draw_line((x, 285), (x + 280, 285), color=(0.3, 0.3, 0.3))
                page.draw_circle((x + 140, 285), 40, color=(0, 0, 1))
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.subset_fonts()
    doc.save(str(out), garbage=4, deflate=True)
    doc.close()
    return out


def make_all(out_dir: Path, sample: dict | None = None) -> tuple[Path, Path]:
    sample = sample or load_sample()
    return make_pptx(sample, out_dir / "sample_report.pptx"), make_pdf(sample, out_dir / "sample_report.pdf")


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("demo")
    for p in make_all(target):
        print(p)

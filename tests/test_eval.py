from pathlib import Path

from openpyxl import load_workbook

from ip_finder.config import Config
from ip_finder.eval import GOLD_HEADERS, evaluate, import_feedback, interpret, make_gold_sheet, train_classifier
from ip_finder.eval.score import load_gold
from ip_finder.history import History
from ip_finder.pipeline import extract


def test_gold_sheet_and_scoring(sample_dir, cfg, fake_client, tmp_path):
    gold = make_gold_sheet(sample_dir / "sample_report.pptx", cfg, tmp_path / "g.xlsx", fake_client)
    wb = load_workbook(str(gold))
    assert wb.sheetnames == ["라벨링", "안내", "메타"]
    header = [c.value for c in wb["라벨링"][1]]
    assert header == GOLD_HEADERS
    rows = [r for r in wb["라벨링"].iter_rows(min_row=2, values_only=True) if r[0]]
    assert rows and all(r[12] in ("A", "B", "C", "D", "X") for r in rows)   # 정답 열이 모델 라벨로 미리 채워짐
    assert all(isinstance(r[10], (int, float)) and isinstance(r[11], (int, float)) for r in rows)   # 점수·가중치
    stored = evaluate(gold, cfg, rerun=False)
    assert stored["raw"]["accuracy"] == 1.0
    rerun = evaluate(gold, cfg, rerun=True, client=fake_client)
    assert rerun["rows"] == len(rows) and rerun["raw"]["accuracy"] == 1.0
    assert set(rerun["final"]["per_grade"]) == {"A", "B", "C", "D", "X"}
    assert rerun["list"]["precision_nonx"] == 1.0 and 0 < rerun["list"]["recall"] <= 1.0 and len(rerun["sweep"]) == 6
    assert stored["list"]["list"] > 0


def test_train_and_use_classifier(sample_dir, cfg, cfg_hybrid, fake_client, fake_embedder, tmp_path):
    gold = make_gold_sheet(sample_dir / "sample_report.pptx", cfg, tmp_path / "g.xlsx", fake_client)
    out = tmp_path / "clf.npz"
    cfg_hybrid.embed_model = "fake-embed"
    stats = train_classifier([gold], cfg_hybrid, out, fake_embedder)
    assert out.exists() and stats["n"] >= 20 and 0.0 <= stats["cv_accuracy"] <= 1.0 and stats["train_accuracy"] > 0.8
    cfg_hybrid.classifier_path = str(out)
    _, res = extract(sample_dir / "sample_report.pptx", cfg_hybrid, tmp_path, fake_client, embedder=fake_embedder)
    assert res.info.classifier.startswith("학습") and res.info.threshold == 0.25
    rep = evaluate(gold, cfg_hybrid, rerun=True, client=fake_client, embedder=fake_embedder)
    assert rep["rows"] == stats["n"] and "hybrid" in rep["source"]


def test_interpret_feedback_values():
    assert interpret("O", "A") == "A" and interpret("o", "C") == "C" and interpret("맞음", "X?") == "X"
    assert interpret("X", "A") == "X" and interpret("틀림", "B") == "X"
    assert interpret("c", "A") == "C" and interpret("", "A") is None and interpret("??", "A") is None


def test_feedback_import_builds_and_merges_gold(sample_dir, cfg, fake_client, tmp_path):
    cfg.history_path = str(tmp_path / "hist")
    out, res = extract(sample_dir / "sample_report.pptx", cfg, tmp_path, fake_client)
    wb = load_workbook(str(out))
    ws = wb["검토 목록"]
    ws.cell(row=2, column=12).value = "O"            # 첫 행 맞음
    ws.cell(row=3, column=12).value = "X"            # 둘째 행 틀림
    ws.cell(row=4, column=12).value = "c"            # 셋째 행 등급 C
    ws.cell(row=4, column=13).value = "원인 설명임"
    ws.cell(row=5, column=12).value = "???"          # 해석 불가 → 건너뜀
    hold = wb["보류"]
    if hold.max_row > 1:
        hold.cell(row=2, column=12).value = "A"      # 보류 문장을 A로 살림
    wb.save(str(out))
    h = History(cfg.history_dir())
    rec = import_feedback(out, history=h)
    gold = Path(rec["gold"])
    assert gold.name == "sample_report_gold.xlsx" and rec["skipped"] == 1 and rec["added"] == rec["feedback_rows"] >= 3
    rows = load_gold(gold)
    first = [r for r in res.labeled if r.label in ("A", "B", "C", "D") and not r.hold_reason]
    by_text = {r["문장"]: r for r in rows}
    assert by_text[first[0].text]["정답"] == first[0].label and by_text[first[1].text]["정답"] == "X" and by_text[first[2].text]["정답"] == "C"
    assert by_text[first[2].text]["메모"] == "원인 설명임" and by_text[first[0].text]["제목"] == first[0].page_title
    assert len(h.feedback_log()) == 1 and h.feedback_log()[0]["corrections"]
    # 두 번째 가져오기: 같은 행은 갱신, 새 행은 추가
    wb = load_workbook(str(out))
    wb["검토 목록"].cell(row=2, column=12).value = "B"
    wb["검토 목록"].cell(row=6, column=12).value = "O"
    wb.save(str(out))
    rec2 = import_feedback(out, history=h)
    rows2 = load_gold(gold)
    assert rec2["updated"] >= 3 and rec2["added"] == 1 and len(rows2) == len(rows) + 1
    assert {r["문장"]: r for r in rows2}[first[0].text]["정답"] == "B"
    # 미표시 행을 맞음으로 간주
    rec3 = import_feedback(out, confirm_unmarked=True, history=h)
    assert len(load_gold(gold)) == len(first) - 1 + (1 if hold.max_row > 1 else 0)   # '???' 행은 끝까지 건너뜀
    assert rec3["gold_rows"] >= len(first)

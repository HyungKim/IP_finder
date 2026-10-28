import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from ip_finder import layout_samples as LS
from ip_finder.readers import load_document


@pytest.fixture(scope="module")
def layout_dir(tmp_path_factory):
    """레이아웃 시험용 합성 보고서 5건 (ip_finder.layout_samples). demo/layout 에 같은 파일이 있다."""
    d = tmp_path_factory.mktemp("layout")
    LS.make_all(d)
    return d


def _texts(page):
    return [ln.text for ln in page.lines]


# ------------------------------------------------------------------ 기존 샘플
def test_pptx_reader_skips_notes_and_reads_tables(sample_dir):
    doc = load_document(sample_dir / "sample_report.pptx")
    assert doc.kind == "pptx" and len(doc.pages) == 12
    all_text = " ".join(ln.text for p in doc.pages for ln in p.lines)
    assert "십자형" not in all_text              # 발표자 노트는 읽지 않는다
    assert doc.pages[4].title.startswith("개선안 1") and doc.pages[4].title_source == "placeholder"
    table_lines = [ln for ln in doc.pages[7].lines if ln.source == "table"]
    assert table_lines and "개선 전" in table_lines[0].text
    assert doc.pages[6].has_graphics and all("대외비" in ln.text for ln in doc.pages[6].lines)   # 도면 슬라이드: 바닥글만


def test_pdf_reader_titles(sample_dir):
    doc = load_document(sample_dir / "sample_report.pdf")
    assert doc.kind == "pdf" and len(doc.pages) == 12
    assert doc.pages[2].title == "배경 및 문제점" and doc.pages[2].title_source == "font"
    assert any("저탄성" in ln.text for ln in doc.pages[4].lines)
    assert doc.pages[6].has_graphics
    assert all("대외비" not in ln.text for p in doc.pages for ln in p.lines)      # 모든 쪽에 반복되는 꼬리말은 리더가 지운다
    assert any("반복 머리말·꼬리말 제거" in n for n in doc.notes)


# ------------------------------------------------------------------ PDF 레이아웃
def test_pdf_single_column_report(layout_dir):
    doc = load_document(layout_dir / "layout_01_single_column.pdf")
    assert len(doc.pages) == 9
    texts = [ln.text for p in doc.pages for ln in p.lines]
    assert not any("대외비" in t or t.startswith("- ") for t in texts)          # 머리말·꼬리말·쪽 번호 제거
    cover = doc.pages[0]
    assert cover.title.startswith("SV-700") and _texts(cover)[-1] == "2026. 10. 28."   # 날짜 줄은 앞줄과 합치지 않는다
    p3 = doc.pages[2]
    assert p3.title == "배경 및 문제점" and p3.title_source == "font"
    assert p3.lines[0].text == LS.WRAPPED_BULLET and p3.lines[0].level == 0     # 두 줄로 꺾인 불릿이 한 문장으로
    assert p3.lines[1].text == LS.SUB_BULLET and p3.lines[1].level == 1         # 하위 불릿은 한 단계 아래
    p4 = doc.pages[3]
    assert p4.lines[0].text == "1) 구조 요인" and p4.lines[1].text == "1. " + LS.NUMBERED_ITEM   # 번호 표식은 남긴다
    assert [ln.level for ln in p4.lines] == [0, 0, 0, 0, 0, 0, 0, 1]
    table = [ln for ln in doc.pages[6].lines if ln.source == "table"]
    assert len(table) == 4
    assert table[0].text.startswith("항목: 진동 속도 RMS (24,000 rpm) | 개선 전: 4.5 mm/s | 개선 후: 1.9 mm/s")
    assert _texts(doc.pages[5])[2:4] == ["제어식", "n_cmd = n_set × (1 + k · sgn(f_res − f_cur) · min(|f_res − f_cur| / f_res, 0.06))"]


def test_pdf_two_columns_read_left_then_right(layout_dir):
    doc = load_document(layout_dir / "layout_02_two_columns.pdf")
    assert any("다단" in n for n in doc.notes)
    p1 = _texts(doc.pages[0])
    assert doc.pages[0].title == "배터리 팩 열관리 제어 로직 검토서"
    assert LS.LEFT_COLUMN_FIRST in p1 and LS.RIGHT_COLUMN_FIRST in p1            # 단 안에서 꺾인 줄은 한 문장으로
    assert p1.index("모듈 온도는 알람에만 사용") < p1.index("2. 제안 방식") < p1.index(LS.RIGHT_COLUMN_FIRST)   # 왼쪽 단을 다 읽은 뒤 오른쪽 단
    p2 = _texts(doc.pages[1])
    assert doc.pages[1].title == "3. 제어 로직 상세" and doc.pages[1].title_source == "font"   # 번호 붙은 소제목도 제목
    assert LS.HYPHEN_SENTENCE in p2                                              # 'com-' + 'puted' → 'computed'
    assert (p2.index("4. 개도 명령 전송 및 15초 대기") < p2.index("4. 영문 요약 (Abstract)")
            < p2.index("5. 시뮬레이션 조건") < p2.index("항목: 급속 충전 전류 | 현행: 150 A | 제안: 150 A | 비고: SOC 20→80%"))
    assert sum(1 for ln in doc.pages[1].lines if ln.source == "table") == 4      # 단 아래 전체 폭 표
    p3 = _texts(doc.pages[2])
    heads = [p3.index(h) for h in LS.COLUMN_HEADS]
    assert heads == sorted(heads) and p3.index("7. 결론") > heads[-1]            # 3단 상자는 왼쪽부터, 결론은 그 뒤
    assert "분배 밸브 3개 추가, 모듈 온도 비례 제어" in p3                        # 상자 안에서 꺾인 불릿
    p4 = doc.pages[3].lines
    assert [ln.source for ln in p4] == ["table"] * 3 + ["body"] + ["table"] * 3 + ["body"]   # 표-설명-표-설명 순서


def test_pdf_diagrams_and_image_page(layout_dir):
    doc = load_document(layout_dir / "layout_03_diagrams.pdf")
    p1 = _texts(doc.pages[0])
    for label in LS.BLOCK_LABELS:
        assert label in p1                                                       # 블록도 상자 라벨은 각각 한 줄
    assert any(t.startswith("Confidence 추정기는 카메라") and t.endswith("사이 신뢰도로 변환") for t in p1)   # 옆 상자 라벨이 끼어도 불릿은 안 끊김
    p3 = doc.pages[2]
    assert p3.has_graphics and _texts(p3) == ["그림 3. 시험 차량 전방 카메라 영상 (합성 예시 이미지)"]   # 사진 페이지는 캡션만
    p4 = _texts(doc.pages[3])
    first_step = "1. 카메라 프레임마다 평균 조도와 대비를 계산하고 100 ms 이동 평균을 취한다"
    assert first_step in p4 and p4.index("융합 결과 출력") < p4.index(first_step)   # 순서도를 다 읽은 뒤 오른쪽 설명
    assert _texts(doc.pages[4])[:3] == ["효과 큼 · 적용 쉬움", "신뢰도 기반 가중치 전환", "히스테리시스 0.2 적용"]   # 2×2 상자는 상자 단위


def test_pdf_tables_with_and_without_borders(layout_dir):
    doc = load_document(layout_dir / "layout_04_tables.pdf")
    p2 = [ln.text for ln in doc.pages[1].lines if ln.source == "table"]
    assert len(p2) == 5 and p2[0] == "공정: 세정 | 3월 불량률: 2.4 % | 6월 불량률: 1.6 % | 9월 불량률: 0.9 % | 주요 불량 유형: 잔류 이물, 얼룩"
    assert [ln.source for ln in doc.pages[2].lines] == ["table"] * 4 + ["body"] * 4     # 왼쪽 표, 그다음 오른쪽 설명
    p4 = [ln.text for ln in doc.pages[3].lines if ln.source == "table"]                 # 선 없는 공백 정렬 표
    assert p4 == ["공정: 세정 | 불량률(개선 전): 2.4 % | 불량률(개선 후): 0.9 % | 조치: 세정액 농도 자동 보정",
                  "공정: 도포 | 불량률(개선 전): 3.1 % | 불량률(개선 후): 0.7 % | 조치: 노즐 압력 폐루프 제어",
                  "공정: 경화 | 불량률(개선 전): 1.2 % | 불량률(개선 후): 1.1 % | 조치: 변경 없음",
                  "공정: 검사 | 불량률(개선 전): 0.8 % | 불량률(개선 후): 0.3 % | 조치: 비전 조명 각도 변경"]
    assert any(n.startswith("표 ") for n in doc.notes)


def test_pdf_mixed_language_and_rotation(layout_dir):
    doc = load_document(layout_dir / "layout_05_mixed_rotated.pdf")
    p1 = _texts(doc.pages[0])
    assert LS.ENGLISH_ONLY in p1 and LS.FORMULA in p1                            # 영어·수식 줄은 리더가 버리지 않는다 (전처리 규칙의 몫)
    p2 = doc.pages[1]
    assert p2.title == "3. Diagnostic coverage" and p2.title_source == "font"
    assert _texts(p2)[6:11] == ["DC = 92 %", "SFF = 96 %", "PFH = 3.2E-8 /h", "HFT = 1", "Type B"]   # 짧은 줄은 합치지 않는다
    p3 = doc.pages[2]                                                            # /Rotate 90 페이지, 화면에서는 바로 선 글
    assert p3.title.startswith("5. 모듈 온도 시험 결과") and p3.lines[0].text == LS.ROTATED_LINE
    assert [ln.text for ln in p3.lines if ln.source == "table"][0].startswith("조건: 급속 충전 150 A | 현행 편차: 2.1 ℃")
    p4 = doc.pages[3]                                                            # 옆으로 누운 글은 읽지 않고 센다
    assert _texts(p4) == [LS.SIDEWAYS_CAPTION] and any("세로로 선 문자" in n for n in doc.notes)


def test_pdf_title_with_gap_before_dash_stays_whole(tmp_path):
    """제목이 두 상자("개선안 2" + "- 냉각 유로…")라 큰 간격이 있어도 한 제목. 본문의 '항목 - 값' 줄도 한 줄."""
    font = LS._font()
    doc = pymupdf.open()
    p = doc.new_page(width=842, height=595)
    tw = pymupdf.TextWriter(p.rect)
    tw.append((50, 70), "개선안 2", font=font, fontsize=22)
    tw.append((200, 70), "- 냉각 유로 이중화 구조", font=font, fontsize=22)
    tw.append((70, 130), "• 냉각 유로를 이중 구조로 변경하여 열을 분산시킨다", font=font, fontsize=13)
    tw.append((70, 160), "노즐 압력", font=font, fontsize=13)
    tw.append((180, 160), "-", font=font, fontsize=13)
    tw.append((230, 160), "0.30 MPa 고정에서 폐루프 제어로", font=font, fontsize=13)
    tw.write_text(p)
    path = tmp_path / "dash.pdf"
    doc.save(str(path))
    pg = load_document(path).pages[0]
    assert pg.title == "개선안 2 - 냉각 유로 이중화 구조" and pg.title_source == "font"
    assert _texts(pg) == ["냉각 유로를 이중 구조로 변경하여 열을 분산시킨다", "노즐 압력 - 0.30 MPa 고정에서 폐루프 제어로"]


def test_pdf_small_repeated_title_in_top_band_is_kept(tmp_path):
    """본문(18pt)보다 작은 제목(16pt)이 맨 위 8% 띠 안에 있고 몇 쪽에 반복돼도 제목으로 남는다. 모든 쪽의 문서명 머리말은 지운다."""
    font = LS._font()
    doc = pymupdf.open()
    for k in range(6):
        p = doc.new_page(width=842, height=595)
        tw = pymupdf.TextWriter(p.rect)
        tw.append((600, 20), "PS-200 개선 보고 | 대외비", font=font, fontsize=8)          # 모든 쪽에 반복되는 문서명
        tw.append((40, 42), "3. 개선안 상세" if 1 <= k <= 3 else f"{k + 1}. 다른 절 제목", font=font, fontsize=16)
        tw.append((60, 120), f"• 본문 첫 줄 {k}: 접착제를 저탄성 실리콘계로 변경", font=font, fontsize=18)
        tw.append((60, 154), "• 다이 하부에 응력 완충층을 삽입", font=font, fontsize=18)
        tw.write_text(p)
    path = tmp_path / "band.pdf"
    doc.save(str(path))
    read = load_document(path)
    assert [pg.title for pg in read.pages][1:4] == ["3. 개선안 상세"] * 3 and read.pages[1].title_source == "first-line"
    assert all(_texts(pg)[0].startswith("본문 첫 줄") for pg in read.pages)
    assert all("대외비" not in ln.text for pg in read.pages for ln in pg.lines) and any("대외비" in n for n in read.notes)


def test_pdf_bold_title_same_size_as_body(tmp_path):
    """제목이 본문과 같은 크기지만 굵은 글씨면 제목(bold)."""
    font = LS._font()
    doc = pymupdf.open()
    p = doc.new_page(width=842, height=595)
    tw = pymupdf.TextWriter(p.rect)
    tw.append((50, 80), "Thermal control logic review", font=pymupdf.Font("hebo"), fontsize=14)   # Helvetica-Bold
    tw.append((60, 130), "• 모듈 온도 편차에 비례해 분배 밸브 개도를 바꾼다", font=font, fontsize=14)
    tw.append((60, 160), "• 편차가 0.5℃ 이하로 줄면 기본 분배 비율로 복귀한다", font=font, fontsize=14)
    tw.write_text(p)
    path = tmp_path / "bold.pdf"
    doc.save(str(path))
    pg = load_document(path).pages[0]
    assert pg.title == "Thermal control logic review" and pg.title_source == "bold"
    assert len(pg.lines) == 2


def test_pdf_empty_and_image_only_pages(tmp_path):
    doc = pymupdf.open()
    doc.new_page(width=842, height=595)
    page = doc.new_page(width=842, height=595)
    pix = pymupdf.Pixmap(pymupdf.csGRAY, pymupdf.IRect(0, 0, 40, 30), False)
    pix.clear_with(120)
    page.insert_image(pymupdf.Rect(50, 50, 790, 540), pixmap=pix)
    p = tmp_path / "empty.pdf"
    doc.save(str(p))
    read = load_document(p)
    assert [pg.title_source for pg in read.pages] == ["없음", "없음"]
    assert read.pages[0].lines == [] and not read.pages[0].has_graphics
    assert read.pages[1].lines == [] and read.pages[1].has_graphics


# ------------------------------------------------------------------ PPTX
def test_pptx_title_by_font_size_when_no_placeholder(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])            # 빈 레이아웃: 제목 플레이스홀더 없음
    small = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(0.5))
    small.text_frame.text = "대외비"
    small.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    big = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(8), Inches(0.8))
    big.text_frame.text = "개선안 3: 냉각 유로 이중화"
    big.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    body = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(8), Inches(3))
    body.text_frame.text = "냉각 유로를 이중 구조로 변경하여 열을 분산"
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    p = tmp_path / "t.pptx"
    prs.save(str(p))
    doc = load_document(p)
    pg = doc.pages[0]
    assert pg.title == "개선안 3: 냉각 유로 이중화" and pg.title_source == "font"
    assert [ln.text for ln in pg.lines] == ["대외비", "냉각 유로를 이중 구조로 변경하여 열을 분산"]


def test_pptx_group_shapes_read_in_visual_order(tmp_path):
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    head = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    head.text_frame.text = "융합 제어 구조"
    head.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    group = slide.shapes.add_group_shape()
    for i, label in enumerate(("Radar 전처리", "Confidence 추정기", "가중치 제어기")):
        box = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6 + i * 3.0), Inches(1.5), Inches(2.4), Inches(0.9))
        box.text_frame.text = label
    right = slide.shapes.add_textbox(Inches(5.2), Inches(3.0), Inches(4.3), Inches(1.5))      # XML 순서: 오른쪽을 먼저 추가
    right.text_frame.text = "오른쪽 설명: 야간에는 레이더 가중치를 올린다"
    left = slide.shapes.add_textbox(Inches(0.5), Inches(3.0), Inches(4.3), Inches(1.5))
    left.text_frame.text = "왼쪽 설명: 카메라 신뢰도를 추정한다"
    p = tmp_path / "g.pptx"
    prs.save(str(p))
    pg = load_document(p).pages[0]
    assert pg.title == "융합 제어 구조"
    assert [ln.text for ln in pg.lines] == ["Radar 전처리", "Confidence 추정기", "가중치 제어기",
                                            "왼쪽 설명: 카메라 신뢰도를 추정한다", "오른쪽 설명: 야간에는 레이더 가중치를 올린다"]
    assert pg.has_graphics


def test_pptx_chart_title_becomes_diagram_line(tmp_path):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "야간 오검출률 비교"
    data = CategoryChartData()
    data.categories = ["기존", "제안"]
    data.add_series("오검출률(%)", (8.4, 7.4))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(8), Inches(4.5), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "야간 시나리오 오검출률"
    p = tmp_path / "c.pptx"
    prs.save(str(p))
    doc = load_document(p)
    pg = doc.pages[0]
    assert [(ln.text, ln.source) for ln in pg.lines] == [("야간 시나리오 오검출률", "diagram")] and pg.has_graphics
    assert any("차트" in n for n in doc.notes)


def test_pptx_smartart_points_become_diagram_lines():
    from lxml import etree

    from ip_finder.readers.pptx_reader import A, DGM, DIAGRAM_URI, P, R, _smartart_lines

    data = (f'<dgm:dataModel xmlns:dgm="{DGM[1:-1]}" xmlns:a="{A[1:-1]}"><dgm:ptLst>'
            '<dgm:pt modelId="1"><dgm:t><a:p><a:r><a:t>신뢰도 추정</a:t></a:r><a:br/><a:r><a:t>가중치 전환</a:t></a:r></a:p></dgm:t></dgm:pt>'
            '<dgm:pt modelId="2" type="sibTrans"><dgm:t><a:p><a:r><a:t>무시</a:t></a:r></a:p></dgm:t></dgm:pt>'
            '<dgm:pt modelId="3"><dgm:t><a:p><a:r><a:t>• 융합 출력</a:t></a:r></a:p></dgm:t></dgm:pt>'
            '</dgm:ptLst></dgm:dataModel>')
    frame = (f'<p:graphicFrame xmlns:p="{P[1:-1]}" xmlns:a="{A[1:-1]}" xmlns:dgm="{DGM[1:-1]}" xmlns:r="{R[1:-1]}">'
             f'<a:graphic><a:graphicData uri="{DIAGRAM_URI}"><dgm:relIds r:dm="rId7"/></a:graphicData></a:graphic></p:graphicFrame>')

    class Part:
        blob = data.encode("utf-8")

    class SlidePart:
        def related_part(self, rid):
            assert rid == "rId7"
            return Part()

    class Slide:
        part = SlidePart()

    class Shape:
        _element = etree.fromstring(frame)

    lines = _smartart_lines(Slide(), Shape())
    assert [(ln.text, ln.source) for ln in lines] == [("신뢰도 추정 가중치 전환", "diagram"), ("융합 출력", "diagram")]   # 줄바꿈은 공백, 전환 점은 무시

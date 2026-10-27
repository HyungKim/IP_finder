from openpyxl import load_workbook

from ip_finder.config import Config
from ip_finder.pipeline import extract, find_inputs, output_path
from ip_finder.post import in_review_list
from tests.conftest import FakeEmbedder


def _rows(ws):
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]


def test_extract_pptx_and_pdf_llm_mode(sample_dir, cfg, fake_client, tmp_path):
    for name in ("sample_report.pptx", "sample_report.pdf"):
        out, res = extract(sample_dir / name, cfg, tmp_path, fake_client)
        assert out == output_path(sample_dir / name, tmp_path) and out.exists()
        wb = load_workbook(str(out))
        assert wb.sheetnames == ["검토 목록", "보류", "2차 제외", "페이지", "변경 사항", "제외 페이지", "이미지 페이지", "1차 제외 (참고)", "실행 정보"]
        review = _rows(wb["검토 목록"])
        assert review and all(r[4] in ("A", "B") for r in review)                                  # 기본 추출 등급 A·B
        assert [r[0] for r in review] == list(range(1, len(review) + 1))
        pages = [r[2] for r in review]
        assert pages == sorted(pages)                        # 페이지 순서
        assert any(r[1] == "개선안" and r[4] == "A" for r in review)
        assert all(r[6] for r in review) and all(r[11] == "LLM" for r in review)
        assert all(r[9] == "-" and r[10] in ("A", "B") for r in review)                           # llm 모드: 1차 없음, 2차 = LLM
        assert all(isinstance(r[7], float) and r[7] >= 0.5 and 0 < r[8] <= 1 for r in review)   # 점수·가중치
        assert all(r[13] is None for r in review)                                                 # 피드백 열 비어 있음
        page_rows = _rows(wb["페이지"])
        assert len(page_rows) == 12 and {r[0] for r in page_rows} == set(range(1, 13))
        assert next(r for r in page_rows if r[0] == 5)[2] == "높음" and next(r for r in page_rows if r[0] == 5)[3] == 1.0
        assert _rows(wb["2차 제외"]) and all(r[4] in ("C", "D", "X", "X?") for r in _rows(wb["2차 제외"]))   # llm 모드의 C·D·X 는 2차 제외 시트
        assert any(r[4] == "C" and r[10] == "A" for r in _rows(wb["2차 제외"]))                   # LLM 은 A 였지만 제목 보정으로 C → 제외
        assert any(ll.label == "C" and ll.stage_out == 2 for ll in res.labeled)
        excluded = {r[0]: r[2] for r in _rows(wb["제외 페이지"])}
        assert 2 in excluded and "목차" in excluded[2]
        assert [r[0] for r in _rows(wb["이미지 페이지"])] == [7]
        info = dict(_rows(wb["실행 정보"]))
        assert info["LLM 모델"] == "fake-model" and info["판정 모드"] == "llm" and info["판정 페이지" if "판정 페이지" in info else "라벨링 페이지"] == 7
        assert res.info.llm_calls == 7 and res.info.cache_hits == 0 and res.info.llm_lines == len(res.labeled)
        assert any("(페이지 유형:" in u for u in fake_client.inputs)   # 힌트가 들어감


def test_title_fix_and_safety_net_in_pipeline(sample_dir, cfg, fake_client, tmp_path):
    out, res = extract(sample_dir / "sample_report.pptx", cfg, tmp_path, fake_client)
    fixed = [ll for ll in res.labeled if "보정(A→C)" in ll.notes]
    assert fixed and all(ll.page_title == "원인 분석" for ll in fixed)
    held = [ll for ll in res.labeled if ll.label == "X?"]
    assert any("불량률" in ll.text for ll in held)
    assert "십자형" not in " ".join(ll.text for ll in res.labeled)


def test_fast_mode_uses_only_classifier(sample_dir, cfg_cascade, fake_client, fake_embedder, tmp_path):
    cfg_cascade.mode = "fast"
    out, res = extract(sample_dir / "sample_report.pptx", cfg_cascade, tmp_path, fake_client, embedder=fake_embedder)
    assert res.info.mode == "fast" and fake_client.calls == 0 and res.info.llm_calls == 0
    assert res.info.clf_lines == len(res.labeled) and all(ll.source == "clf" for ll in res.labeled)
    assert all(ll.clf_conf is not None for ll in res.labeled)
    info = dict(_rows(load_workbook(str(out))["실행 정보"]))
    assert info["판정 모드"] == "fast" and info["분류기"].startswith("제로샷")


def test_cascade_classifier_first_then_llm_on_candidates(sample_dir, cfg_cascade, fake_client, fake_embedder, tmp_path):
    out, res = extract(sample_dir / "sample_report.pptx", cfg_cascade, tmp_path, fake_client, embedder=fake_embedder)
    assert res.info.mode == "cascade" and res.info.target_grades == "A,B"
    out1 = [ll for ll in res.labeled if ll.clf_label not in ("A", "B")]
    cands = [ll for ll in res.labeled if ll.clf_label in ("A", "B")]
    assert out1 and cands
    assert all(ll.source == "clf" and ll.llm_label == "" and ll.stage_out == 1 for ll in out1)        # 1차 제외(C·D·X)는 LLM 에 안 감
    assert all(ll.llm_label != "" and ll.source == "llm" for ll in cands)                               # A·B 후보 전부 LLM 2차
    assert res.info.llm_lines == len(cands) and res.info.clf_lines == len(out1)
    assert 1 <= res.info.llm_calls <= 5                                                                  # 페이지 유형별 묶음
    sent = " ".join(fake_client.inputs)
    assert all(ll.text not in sent for ll in out1)
    rejected = [ll for ll in cands if ll.llm_label not in ("A", "B")]
    assert res.info.llm_rejected == len(rejected) and all(ll.stage_out == 2 and "LLM 2차 제외(" in " ".join(ll.notes) for ll in rejected)
    passed = [ll for ll in cands if ll.stage_out == 0]
    assert passed and all(ll.label in ("A", "B") for ll in passed)
    agree = [ll for ll in passed if ll.llm_label == ll.clf_label]
    assert agree and all(ll.confidence == 1.0 for ll in agree)
    wb = load_workbook(str(out))
    second = _rows(wb["2차 제외"])
    assert len(second) == len(rejected) and all(r[9] in ("A", "B") and r[10] in ("C", "D", "X") for r in second)   # 1차 통과, 2차에서 제외
    assert not any(r[10] in ("C", "D", "X") for r in _rows(wb["보류"]))                               # 보류에는 2차 제외가 없음
    first_out = _rows(wb["1차 제외 (참고)"])
    assert len(first_out) == sum(1 for ll in out1 if ll.label != "X?") and all(r[9] in ("C", "D", "X") and r[10] == "-" for r in first_out)
    review = _rows(wb["검토 목록"])
    assert review and all(r[4] in ("A", "B") and r[9] in ("A", "B") and r[10] in ("A", "B") for r in review)
    info = dict(_rows(wb["실행 정보"]))
    assert info["단계별 문장 수"].startswith(f"전체 {len(res.labeled)} → 1차 제외") and info["추출 등급"] == "A,B"
    # 등급을 넓히면 C·D 도 후보가 된다
    cfg_cascade.target_grades = "A,B,C,D"
    _, res4 = extract(sample_dir / "sample_report.pptx", cfg_cascade, tmp_path, type(fake_client)(), embedder=fake_embedder)
    assert res4.info.llm_lines > res.info.llm_lines and any(ll.label == "C" and in_review_list(ll) for ll in res4.labeled)


def test_cascade_falls_back_to_llm_without_embedder(sample_dir, cfg_cascade, fake_client, tmp_path):
    _, res = extract(sample_dir / "sample_report.pptx", cfg_cascade, tmp_path, fake_client, embedder=FakeEmbedder(available=False))
    assert res.info.mode == "llm" and res.info.warnings and "llm 모드로 대체" in res.info.warnings[0]
    assert fake_client.calls == 7


def test_cache_reuses_labels_and_vectors(sample_dir, fake_client, fake_embedder, tmp_path):
    cfg = Config(use_cache=True, cache_path=str(tmp_path / "cache.sqlite"), mode="llm", history_path=str(tmp_path / "hist"), gold_paths=[str(tmp_path / "gold")])
    _, first = extract(sample_dir / "sample_report.pptx", cfg, tmp_path, fake_client)
    _, second = extract(sample_dir / "sample_report.pptx", cfg, tmp_path, fake_client)
    assert first.info.llm_calls == 7 and first.info.cache_hits == 0
    assert second.info.llm_calls == 0 and second.info.cache_hits == 7
    assert [ll.label for ll in first.labeled] == [ll.label for ll in second.labeled]
    raw = (tmp_path / "cache.sqlite").read_bytes()
    assert "저탄성".encode("utf-8") not in raw                # 캐시에 원문 미저장


def test_find_inputs(sample_dir):
    files = find_inputs([sample_dir])
    assert {f.suffix for f in files} == {".pdf", ".pptx"}


def test_page_weight_moves_low_pages_to_hold(sample_dir, cfg, fake_client, tmp_path):
    cfg.target_grades = "A,B,C,D"                                            # 원인 분석 페이지는 C 문장이라 등급을 넓혀 시험
    cfg.low_relevance_patterns, cfg.low_relevance_weight = ["원인"], 0.3     # 원인 분석 페이지 → 관련성 낮음 0.3 → 점수 미달
    out, res = extract(sample_dir / "sample_report.pptx", cfg, tmp_path, fake_client)
    held = [ll for ll in res.labeled if ll.page_title == "원인 분석" and ll.label in ("A", "B", "C", "D")]
    assert held and all(ll.hold_reason.startswith("점수") and ll.page_weight == 0.3 for ll in held)
    assert res.info.held_lines == len(held) and res.info.pages_low == 1
    wb = load_workbook(str(out))
    assert not any(r[2] == 4 for r in _rows(wb["검토 목록"]))
    hold = _rows(wb["보류"])
    assert any(r[2] == 4 and str(r[15]).startswith("점수") for r in hold)
    assert next(r for r in _rows(wb["페이지"]) if r[0] == 4)[2] == "낮음"
    cfg.page_weight_enabled = False                                          # 스위치 끔 → 전부 1.0, 보류 없음
    _, res_off = extract(sample_dir / "sample_report.pptx", cfg, tmp_path, type(fake_client)())
    assert all(ll.page_weight == 1.0 for ll in res_off.labeled) and res_off.info.held_lines == 0
    assert all(p.relevance == "높음" and p.weight_reason == "가중치 끔" for p in res_off.doc.pages)
    cfg.page_weight_enabled, cfg.low_relevance_weight = True, 0.0           # 낮음 가중치 0 → 페이지 자체를 판정하지 않음
    _, res0 = extract(sample_dir / "sample_report.pptx", cfg, tmp_path, type(fake_client)())
    assert not any(ll.page_title == "원인 분석" for ll in res0.labeled)
    assert any(p.excluded_reason and "가중치 0" in p.excluded_reason for p in res0.doc.pages)


def test_history_diff_between_runs(sample_dir, fake_client, tmp_path):
    cfg = Config(use_cache=False, mode="llm", history_path=str(tmp_path / "hist"), gold_paths=[str(tmp_path / "gold")])
    out1, first = extract(sample_dir / "sample_report.pptx", cfg, tmp_path, fake_client)
    assert first.info.run_id and not first.info.prev_run_id and first.changes == []
    raw = b"".join(p.read_bytes() for p in (tmp_path / "hist" / "runs").glob("*.json"))
    assert "저탄성".encode("utf-8") not in raw and "원인 분석".encode("utf-8") not in raw   # 이력에 원문·제목 없음
    cfg.accept_score = 0.95                      # 분류기 없는 llm 모드 신뢰도 0.9 → 전부 보류
    out2, second = extract(sample_dir / "sample_report.pptx", cfg, tmp_path, type(fake_client)())
    assert second.info.prev_run_id == first.info.run_id
    assert any("목록 점수 기준" in s for s in second.info.settings_diff)
    assert second.info.changes["left"] == sum(1 for ll in first.labeled if in_review_list(ll))
    assert second.changes and all(r["kind"] == "목록 이탈" and r["text"] for r in second.changes)
    wb = load_workbook(str(out2))
    rows = [r for r in wb["변경 사항"].iter_rows(min_row=2, values_only=True)]
    assert rows[0][0] == "기준" and any(r[0] == "설정" and "목록 점수 기준" in str(r[3]) for r in rows)
    assert sum(1 for r in rows if r[0] == "목록 이탈") == second.info.changes["left"]
    info = dict(_rows(wb["실행 정보"]))
    assert info["직전 실행 ID"] == first.info.run_id
